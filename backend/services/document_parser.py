import os
import tempfile
import pymupdf
from docx import Document
from pptx import Presentation

class DocumentParser:
    @staticmethod
    def parse_pdf(file_path: str) -> str:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        from PIL import Image
        import io
        text = ""
        try:
            doc = pymupdf.open(file_path)
            for page in doc:
                page_text = page.get_text().strip()
                if not page_text or len(page_text) < 150:
                    # Fallback to OCR if the page has very little embedded text (e.g. just a CamScanner watermark)
                    pix = page.get_pixmap(dpi=300)
                    img = Image.open(io.BytesIO(pix.tobytes()))
                    page_text = pytesseract.image_to_string(img).strip()
                text += page_text + "\n"
        except Exception as e:
            print(f"Error parsing PDF: {e}")
        return text

    @staticmethod
    def parse_docx(file_path: str) -> str:
        text = ""
        try:
            doc = Document(file_path)
            for para in doc.paragraphs:
                text += para.text + "\n"
        except Exception as e:
            print(f"Error parsing DOCX: {e}")
        return text

    @staticmethod
    def parse_pptx(file_path: str) -> str:
        text = ""
        try:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
        except Exception as e:
            print(f"Error parsing PPTX: {e}")
        return text

    @staticmethod
    def parse_image(file_path: str) -> str:
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        from PIL import Image
        text = ""
        try:
            img = Image.open(file_path)
            text = pytesseract.image_to_string(img).strip()
        except Exception as e:
            print(f"Error parsing Image: {e}")
        return text

    @staticmethod
    def parse_file(file_path: str, file_type: str) -> str:
        if file_type == 'application/pdf':
            return DocumentParser.parse_pdf(file_path)
        elif file_type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
            return DocumentParser.parse_docx(file_path)
        elif file_type == 'application/vnd.openxmlformats-officedocument.presentationml.presentation':
            return DocumentParser.parse_pptx(file_path)
        elif file_type.startswith('image/'):
            return DocumentParser.parse_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {file_type}")
